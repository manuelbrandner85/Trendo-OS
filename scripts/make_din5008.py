#!/usr/bin/env python3
"""Erzeugt eigene DIN-5008-Briefvorlagen (Form B) als DOCX – eigenes Projekt-Asset, frei nutzbar."""
from docx import Document
from docx.shared import Mm, Pt
import pathlib
OUT = pathlib.Path(__file__).resolve().parent.parent / "src/Trendo.App/Assets/templates"
def letter(name, business):
    d = Document(); s = d.sections[0]
    s.page_width, s.page_height = Mm(210), Mm(297)
    s.left_margin, s.right_margin, s.top_margin, s.bottom_margin = Mm(25), Mm(20), Mm(27), Mm(20)
    st = d.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(11)
    p = d.add_paragraph("Vorname Nachname · Straße Nr. · PLZ Ort"); p.runs[0].font.size = Pt(8)  # Rücksendeangabe (Form B: ab 45 mm)
    p.paragraph_format.space_before = Mm(18)
    for line in ["Firma / Frau / Herrn", "Vorname Nachname", "Straße Nr.", "PLZ Ort"]: d.add_paragraph(line).paragraph_format.space_after = Pt(0)
    d.add_paragraph().paragraph_format.space_before = Mm(20)
    if business:
        t = d.add_table(rows=2, cols=4)
        for i, h in enumerate(["Ihr Zeichen", "Unser Zeichen", "Telefon", "Datum"]): t.cell(0, i).text = h
        t.cell(1, 3).text = "TT.MM.JJJJ"
    else:
        d.add_paragraph("Ort, TT.MM.JJJJ").alignment = 2
    d.add_paragraph().add_run("Betreff").bold = True
    d.add_paragraph("Sehr geehrte Damen und Herren,")
    d.add_paragraph("Text des Briefes.")
    d.add_paragraph("Mit freundlichen Grüßen")
    d.add_paragraph("\n\nVorname Nachname")
    d.core_properties.title = name; d.core_properties.author = "Trendo OS"
    d.save(OUT / f"{name}.docx")
letter("DIN5008-Geschaeftsbrief-FormB", True); letter("DIN5008-Privatbrief-FormB", False); print("ok")
